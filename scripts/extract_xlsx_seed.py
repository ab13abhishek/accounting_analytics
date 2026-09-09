"""
One-time extraction script: reads the source practice workbook
(Bookkeeping_50_Cases_EndToEnd.xlsx) and converts it into clean JSON seed
files under app/seed_data/. The FastAPI app never touches the .xlsx file
directly (and does not depend on openpyxl at runtime) -- it only reads the
JSON produced here.

Usage:
    python scripts/extract_xlsx_seed.py "C:\\path\\to\\Bookkeeping_50_Cases_EndToEnd.xlsx"

Also writes app/seed_data/expected_results.json, a snapshot of the
workbook's own computed Ledger / Trial Balance / P&L figures for each case.
These are NOT used by the running app -- they exist purely so
tests/test_accounting_engine.py can prove the engine's computed numbers
reconcile exactly to the source workbook (the same "trace it back" check a
CFA / bookkeeping review would perform).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

SEED_DIR = Path(__file__).resolve().parent.parent / "app" / "seed_data"


def parse_index(ws) -> list[dict]:
    companies = []
    header_row = None
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[0] == "Case #":
            header_row = row
            continue
        if header_row is None:
            continue
        case_no = row[0]
        if not isinstance(case_no, int):
            continue
        companies.append(
            {
                "case_no": case_no,
                "category": row[1],
                "industry": row[2],
                "business_name": row[3],
                "journal_tab": row[4],
                "ledger_tab": row[5],
                "tb_tab": row[6],
                "pnl_tab": row[7],
                "journal_rows": row[8],
            }
        )
    return companies


def parse_journal(ws) -> list[dict]:
    lines = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            if row[0] == "Date":
                header_seen = True
            continue
        date, txn_id, description, acct_code, acct_name, gst_code, debit, credit = row[:8]
        if date is None or txn_id is None:
            continue
        if isinstance(acct_code, str) and "check" in (description or "").lower():
            continue
        lines.append(
            {
                "date": date.isoformat() if hasattr(date, "isoformat") else str(date),
                "txn_id": int(txn_id) if isinstance(txn_id, (int, float)) else txn_id,
                "description": description,
                "account_code": acct_code,
                "account_name": acct_name,
                "gst_code": gst_code,
                "debit": round(float(debit), 2) if debit else 0.0,
                "credit": round(float(credit), 2) if credit else 0.0,
            }
        )
    return lines


def parse_ledger(ws) -> list[dict]:
    rows = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            if row[0] == "Account Code":
                header_seen = True
            continue
        code, name, normal, td, tc, cb, btype = row[:7]
        if code is None:
            continue
        rows.append(
            {
                "account_code": code,
                "account_name": name,
                "normal_balance": normal,
                "total_debits": round(float(td or 0), 2),
                "total_credits": round(float(tc or 0), 2),
                "closing_balance": round(float(cb or 0), 2),
                "balance_type": btype,
            }
        )
    return rows


def parse_tb(ws) -> list[dict]:
    rows = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not header_seen:
            if row[0] == "Account Code":
                header_seen = True
            continue
        code, name, db, cb = row[:4]
        if code is None:
            continue
        rows.append(
            {
                "account_code": code,
                "account_name": name,
                "debit_balance": round(float(db or 0), 2),
                "credit_balance": round(float(cb or 0), 2),
            }
        )
    return rows


def parse_pnl(ws) -> dict:
    income, expenses = [], []
    section = None
    total_income = total_expenses = net_profit = None
    for row in ws.iter_rows(values_only=True):
        a, b, c = (row + (None, None, None))[:3]
        if a == "INCOME":
            section = "income"
            continue
        if a == "EXPENSES":
            section = "expenses"
            continue
        if a == "Account Code":
            continue
        if b == "TOTAL INCOME":
            total_income = round(float(c or 0), 2)
            continue
        if b == "TOTAL EXPENSES":
            total_expenses = round(float(c or 0), 2)
            continue
        if b and "NET PROFIT" in str(b).upper():
            net_profit = round(float(c or 0), 2)
            continue
        if a and section in ("income", "expenses") and c is not None:
            target = income if section == "income" else expenses
            target.append({"account_code": a, "account_name": b, "amount": round(float(c or 0), 2)})
    return {
        "income": income,
        "expenses": expenses,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "net_profit": net_profit,
    }


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    companies = parse_index(wb["INDEX"])

    chart_of_accounts: dict[str, dict] = {}
    all_journals: dict[int, list[dict]] = {}
    expected_results: dict[int, dict] = {}

    for c in companies:
        n = c["case_no"]
        tag = f"C{n:02d}"

        journal = parse_journal(wb[f"{tag}_Journal"])
        ledger = parse_ledger(wb[f"{tag}_Ledger"])
        tb = parse_tb(wb[f"{tag}_TB"])
        pnl = parse_pnl(wb[f"{tag}_PnL"])

        all_journals[n] = journal
        expected_results[n] = {"ledger": ledger, "trial_balance": tb, "pnl": pnl}

        for row in ledger:
            code = row["account_code"]
            if code not in chart_of_accounts:
                acct_type = (
                    "ASSET"
                    if code.startswith("1-")
                    else "LIABILITY"
                    if code.startswith("2-")
                    else "EQUITY"
                    if code.startswith("3-")
                    else "INCOME"
                    if code.startswith("4-")
                    else "COGS"
                    if code.startswith("5-")
                    else "EXPENSE"
                )
                chart_of_accounts[code] = {
                    "code": code,
                    "name": row["account_name"],
                    "account_type": acct_type,
                    "normal_balance": row["normal_balance"].upper(),
                }

    SEED_DIR.mkdir(parents=True, exist_ok=True)

    (SEED_DIR / "chart_of_accounts.json").write_text(
        json.dumps(sorted(chart_of_accounts.values(), key=lambda a: a["code"]), indent=2),
        encoding="utf-8",
    )
    (SEED_DIR / "companies.json").write_text(json.dumps(companies, indent=2), encoding="utf-8")
    (SEED_DIR / "journal_entries.json").write_text(
        json.dumps(all_journals, indent=2), encoding="utf-8"
    )
    (SEED_DIR / "expected_results.json").write_text(
        json.dumps(expected_results, indent=2), encoding="utf-8"
    )

    print(f"Extracted {len(companies)} companies")
    print(f"Chart of accounts: {len(chart_of_accounts)} accounts")
    print(f"Total journal lines: {sum(len(v) for v in all_journals.values())}")
    print(f"Wrote seed files to {SEED_DIR}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
